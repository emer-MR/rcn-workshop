"""Wtyczki -- REST API: lista, uruchomienie na koszyku transakcji.

Wtyczki to pojedyncze pliki `.py` w `data/plugins/` (globalne, nie per-workspace).
Uruchomienie: POST /api/workspaces/{wid}/plugins/{plugin_id} z tym samym body co
Σ Analiza (`{id_rcn_in: [...]}`) -- wiersze ładuje wspólny `_load_rows`, wynik
(JSON: tables/charts/text/...) renderuje generyczny renderer w workspace.js.
"""
from __future__ import annotations

import json
import logging
import re
import uuid

from fastapi import APIRouter, Depends, File, Form, HTTPException, Response, UploadFile
from pydantic import BaseModel, Field

from app.analysis import MAX_IDS, AnalysisRequest, _load_rows
from app.auth import require_admin, require_auth
from app.config import settings
from app.plugin_loader import (
    DISABLED_SUFFIX,
    PLUGIN_ID_RE,
    _load_plugin_file,
    get_plugin,
    get_registry,
    scan_plugins,
)
from app.workspaces import _poi_sqlite, _require_workspace, _workspace_db
from rcn_core.ingest import open_workspace

log = logging.getLogger(__name__)
router = APIRouter(prefix="/api", tags=["plugins"])

# Wtyczka to pojedynczy plik .py -- 1 MB to i tak o rzad wielkosci za duzo.
MAX_PLUGIN_BYTES = 1 * 1024 * 1024


@router.get("/plugins")
def list_plugins(_: str = Depends(require_auth)) -> list[dict]:
    return [rec.to_api() for rec in get_registry().values()]


@router.post("/plugins/rescan")
def rescan_plugins(_: str = Depends(require_auth)) -> list[dict]:
    return [rec.to_api() for rec in scan_plugins().values()]


@router.post("/plugins/upload")
async def upload_plugin(
    file: UploadFile = File(...),
    confirm: bool = Form(False),
    overwrite: bool = Form(False),
    _: str = Depends(require_admin),
) -> dict:
    """Upload wtyczki -- dwufazowo: `confirm=false` waliduje (= WYKONUJE kod,
    świadoma decyzja: admin + dialog ostrzegawczy w UI) i zwraca preview metadanych;
    `confirm=true` zapisuje atomowo do `plugins_dir/{PLUGIN_id}.py` (nazwa Z ID,
    NIGDY z filename uploadu). 409 przy kolizji id bez `overwrite`."""
    fn = (file.filename or "").lower()
    if not fn.endswith(".py"):
        raise HTTPException(status_code=400, detail="Plik wtyczki musi mieć rozszerzenie .py")

    pdir = settings.plugins_dir
    pdir.mkdir(parents=True, exist_ok=True)
    tmp = pdir / f".upload-{uuid.uuid4().hex}.py"
    try:
        total = 0
        with open(tmp, "wb") as handle:
            while True:
                chunk = await file.read(64 * 1024)
                if not chunk:
                    break
                total += len(chunk)
                if total > MAX_PLUGIN_BYTES:
                    raise HTTPException(status_code=413, detail="Wtyczka przekracza limit 1 MB")
                handle.write(chunk)
        await file.close()

        rec = _load_plugin_file(tmp)
        if rec.error:
            raise HTTPException(status_code=400, detail=f"Walidacja nieudana: {rec.error}")

        preview = rec.to_api()
        preview["filename"] = f"{rec.plugin_id}.py"
        target = pdir / f"{rec.plugin_id}.py"
        target_disabled = pdir / f"{rec.plugin_id}{DISABLED_SUFFIX}"
        exists = target.exists() or target_disabled.exists()

        if not confirm:
            return {"status": "preview", "plugin": preview, "exists": exists}

        if exists and not overwrite:
            raise HTTPException(
                status_code=409,
                detail=f"Wtyczka '{rec.plugin_id}' już istnieje (wyślij overwrite=true aby nadpisać)",
            )
        # Nadpisanie zdejmuje też wariant wyłączony -- nowa wersja startuje włączona.
        target_disabled.unlink(missing_ok=True)
        tmp.replace(target)
        scan_plugins()
        return {"status": "installed", "plugin": preview}
    finally:
        tmp.unlink(missing_ok=True)


def _plugin_file_pairs(plugin_id: str) -> list[tuple]:
    """Pary (aktywny, wyłączony) plików wtyczki: najpierw ścieżka z rejestru
    (działa też gdy nazwa pliku != PLUGIN['id'], np. wtyczka skopiowana ręcznie),
    potem warianty z konwencji nazwy `{id}.py`. Ścieżka z rejestru jest bezpieczna
    (pochodzi ze skanu plugins_dir); konwencja nazwy wymaga walidacji id
    (ochrona path traversal)."""
    pairs: list[tuple] = []
    rec = get_plugin(plugin_id)
    if rec is not None and rec.path is not None:
        name = rec.path.name
        base = name.removesuffix(DISABLED_SUFFIX) if name.endswith(DISABLED_SUFFIX) \
            else name.removesuffix(".py")
        pairs.append((rec.path.parent / f"{base}.py",
                      rec.path.parent / f"{base}{DISABLED_SUFFIX}"))
    if PLUGIN_ID_RE.fullmatch(plugin_id):
        pdir = settings.plugins_dir
        pair = (pdir / f"{plugin_id}.py", pdir / f"{plugin_id}{DISABLED_SUFFIX}")
        if pair not in pairs:
            pairs.append(pair)
    if not pairs:
        # Ani w rejestrze, ani poprawny id -- nie da się bezpiecznie zbudować ścieżki.
        raise HTTPException(status_code=400, detail="Niepoprawny identyfikator wtyczki")
    return pairs


@router.delete("/plugins/{plugin_id}")
def delete_plugin(plugin_id: str, _: str = Depends(require_admin)) -> dict:
    removed = False
    for pair in _plugin_file_pairs(plugin_id):
        for path in pair:
            if path.exists():
                path.unlink()
                removed = True
    if not removed:
        raise HTTPException(status_code=404, detail=f"Wtyczka '{plugin_id}' nie istnieje")
    scan_plugins()
    return {"status": "deleted", "id": plugin_id}


class EnableRequest(BaseModel):
    enabled: bool


@router.post("/plugins/{plugin_id}/enable")
def enable_plugin(plugin_id: str, body: EnableRequest, _: str = Depends(require_admin)) -> dict:
    pairs = _plugin_file_pairs(plugin_id)
    active, disabled = next(
        (pair for pair in pairs if pair[0].exists() or pair[1].exists()),
        (None, None),
    )
    if active is None:
        raise HTTPException(status_code=404, detail=f"Wtyczka '{plugin_id}' nie istnieje")
    if body.enabled:
        if not active.exists() and disabled.exists():
            disabled.replace(active)
    else:
        if not disabled.exists() and active.exists():
            active.replace(disabled)
    scan_plugins()
    rec = get_plugin(plugin_id)
    return rec.to_api() if rec else {"id": plugin_id, "enabled": body.enabled}


class PluginResultExportRequest(BaseModel):
    """Body generycznego eksportu wyniku wtyczki: dokładnie te tabele (i tekst),
    które wtyczka zwróciła z run() -- frontend odsyła je bez przetwarzania."""
    filename: str = "wynik-wtyczki"
    tables: list = []
    text: str | None = None


_SHEET_FORBIDDEN = re.compile(r"[\\/*?:\[\]]")
_MAX_EXPORT_CELLS = 500_000


@router.post("/plugins/result.xlsx")
def export_plugin_result(body: PluginResultExportRequest, _: str = Depends(require_auth)) -> Response:
    """Wynik wtyczki (tabele JSON z kontraktu renderera) -> skoroszyt XLSX:
    arkusz per tabela (nagłówki pogrubione, freeze panes, dopasowane szerokości)
    + opcjonalny arkusz „Interpretacja" z polem text. Realizuje eksport dla
    KAŻDEJ wtyczki bez dedykowanego kodu."""
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    used_titles: set[str] = set()
    total_cells = 0

    def _sheet_title(raw, idx: int) -> str:
        title = _SHEET_FORBIDDEN.sub(" ", str(raw or "")).strip()[:28] or f"Tabela {idx + 1}"
        base, n = title, 2
        while title in used_titles:
            title = f"{base[:25]} ({n})"
            n += 1
        used_titles.add(title)
        return title

    for idx, table in enumerate(body.tables):
        if not isinstance(table, dict):
            continue
        if isinstance(table.get("columns"), list) and isinstance(table.get("rows"), list):
            cols = [str(c) for c in table["columns"]]
            rows = [r if isinstance(r, list) else [r] for r in table["rows"]]
        else:
            # Tabela dict klucz->wartość (klucz 'title' to nagłówek, nie wiersz).
            cols = ["Wielkość", "Wartość"]
            rows = [[k, v] for k, v in table.items() if k != "title"]
        total_cells += (len(rows) + 1) * max(1, len(cols))
        if total_cells > _MAX_EXPORT_CELLS:
            raise HTTPException(status_code=413, detail="Wynik zbyt duży do eksportu XLSX")
        ws = wb.create_sheet(_sheet_title(table.get("title"), idx))
        ws.append(cols)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = "A2"
        for row in rows:
            ws.append([("" if v is None else v if isinstance(v, (int, float, bool)) else str(v))
                       for v in row])
        for ci in range(1, len(cols) + 1):
            sample = [len(str(r[ci - 1])) for r in rows[:200] if len(r) >= ci and r[ci - 1] is not None]
            width = max([len(cols[ci - 1]), *sample]) if sample else len(cols[ci - 1])
            ws.column_dimensions[get_column_letter(ci)].width = min(max(10, width + 2), 60)

    if body.text:
        ws = wb.create_sheet(_sheet_title("Interpretacja", len(body.tables)))
        for line in str(body.text).split("\n"):
            ws.append([line])
        ws.column_dimensions["A"].width = 120

    if not wb.sheetnames:
        raise HTTPException(status_code=400, detail="Brak tabel do eksportu")

    buf = BytesIO()
    wb.save(buf)
    fname = re.sub(r"[^A-Za-z0-9_.-]+", "-", body.filename).strip("-.") or "wynik-wtyczki"
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}.xlsx"'},
    )


class PluginRunRequest(AnalysisRequest):
    """Body uruchomienia wtyczki: koszyk jak w Σ Analizie + opcjonalne parametry
    przekazywane wtyczce przez ctx.params (np. alpha, data_wyceny).

    `id_rcn_in` BEZ pydantic max_length (inaczej 422 z kryptycznym komunikatem):
    nadmiar przycinamy do MAX_IDS i sygnalizujemy nagłówkiem X-RCN-Capped."""
    id_rcn_in: list[str] = Field(..., min_length=1)
    params: dict = {}


@router.post("/workspaces/{workspace_id}/plugins/{plugin_id}")
def run_plugin(
    workspace_id: str,
    plugin_id: str,
    body: PluginRunRequest,
    _: str = Depends(require_auth),
) -> Response:
    _require_workspace(workspace_id)
    rec = get_plugin(plugin_id)
    if rec is None:
        raise HTTPException(status_code=404, detail=f"Wtyczka '{plugin_id}' nie istnieje")
    if not rec.enabled:
        raise HTTPException(status_code=409, detail=f"Wtyczka '{plugin_id}' jest wyłączona")
    if rec.error:
        raise HTTPException(status_code=400, detail=f"Wtyczka '{plugin_id}' ma błąd: {rec.error}")
    if rec.kind == "export":
        raise HTTPException(status_code=501, detail="Wtyczki kind=export nie są jeszcze obsługiwane")

    ids = list({i for i in body.id_rcn_in if i})
    if not ids:
        raise HTTPException(status_code=400, detail="id_rcn_in must not be empty")
    requested = len(ids)
    if requested > MAX_IDS:
        ids = ids[:MAX_IDS]

    conn = open_workspace(_workspace_db(workspace_id))
    try:
        rows = _load_rows(conn, ids)
    finally:
        conn.close()

    from app.plugin_ctx import PluginCtx
    ctx = PluginCtx(_poi_sqlite(workspace_id), params=body.params)
    try:
        result = rec.run(rows, ctx) if rec.accepts_ctx else rec.run(rows)
    except Exception as exc:
        log.exception("Wtyczka %s: wyjątek podczas run()", plugin_id)
        raise HTTPException(
            status_code=500,
            detail=f"Wtyczka '{plugin_id}' rzuciła wyjątek: {type(exc).__name__}",
        ) from exc
    finally:
        ctx.close()

    if not isinstance(result, dict):
        raise HTTPException(
            status_code=500,
            detail=f"Wtyczka '{plugin_id}' zwróciła {type(result).__name__} zamiast dict",
        )
    try:
        # allow_nan=False -- NaN/inf w wyniku to błąd wtyczki, nie niepoprawny JSON u klienta.
        payload = json.dumps(result, ensure_ascii=False, allow_nan=False)
    except ValueError as exc:
        raise HTTPException(
            status_code=500,
            detail=f"Wtyczka '{plugin_id}' zwróciła wartości NaN/inf (niepoprawny JSON)",
        ) from exc
    headers = {}
    if requested > MAX_IDS:
        # Sygnał przycięcia koszyka do MAX_IDS -- UI pokazuje notkę nad wynikiem.
        headers["X-RCN-Capped"] = f"{MAX_IDS}/{requested}"
    return Response(content=payload, media_type="application/json", headers=headers)
