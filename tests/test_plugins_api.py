"""Testy REST API wtyczek: lista, uruchomienie, upload/zarządzanie, kody błędów
+ regresja odpowiedzi /analysis po refaktorze _load_rows."""
from __future__ import annotations

import sys
import time

import pytest

ECHO_SRC = '''
PLUGIN = {"id": "echo", "name": "Echo", "version": "1.0",
          "author": "test", "kind": "analysis", "description": "zwraca n"}

def run(rows, ctx):
    with_m2 = [r for r in rows if r.get("cena_m2") is not None]
    return {"text": f"n={len(rows)} m2={len(with_m2)}",
            "table": {"n": len(rows), "has_poi": ctx.has_poi() if ctx else None,
                      "alpha": ctx.params.get("alpha") if ctx else None}}
'''


def _write_plugin(tmp_data_dir, filename, src):
    pdir = tmp_data_dir / "plugins"
    pdir.mkdir(exist_ok=True)
    (pdir / filename).write_text(src, encoding="utf-8")


def _make_workspace_with_tx(tmp_data_dir, name="Testowo") -> str:
    """Workspace ze schematem + 4 transakcje (3 lokale z area, 1 grunt bez)."""
    from rcn_core.ingest import open_workspace
    wdir = tmp_data_dir / "workspaces" / name
    wdir.mkdir(parents=True)
    conn = open_workspace(wdir / "workspace.sqlite")
    now = int(time.time())
    conn.execute(
        "INSERT INTO imports(original_filename, stored_filename, tryb, upload_timestamp, status) "
        "VALUES('t.gml','t.gml','snapshot',?,'completed')", (now,))
    rows = [
        ("RCN-1", "2025-01-10", 400000.0, "rynek wtórny", "sprzedaż", "lokal", 50.0, 19.45, 51.75),
        ("RCN-2", "2025-02-15", 450000.0, "rynek wtórny", "sprzedaż", "lokal", 55.0, 19.46, 51.76),
        ("RCN-3", "2025-03-20", 500000.0, "rynek pierwotny", "sprzedaż", "lokal", 60.0, 19.47, 51.77),
        ("RCN-4", "2025-04-25", 300000.0, "rynek wtórny", "sprzedaż", "grunt niezabudowany", None, None, None),
    ]
    for id_rcn, data, cena, rynek, rodzaj_tx, rodzaj_nier, area, lon, lat in rows:
        conn.execute(
            "INSERT INTO transakcje(id_rcn, source_import_id, import_timestamp, data_transakcji, "
            "cena_transakcji_brutto, rodzaj_rynku, rodzaj_transakcji, rodzaj_nieruchomosci) "
            "VALUES(?,1,?,?,?,?,?,?)",
            (id_rcn, now, data, cena, rynek, rodzaj_tx, rodzaj_nier))
        conn.execute(
            "INSERT INTO tx_cache(id_rcn, miejscowosc, adres, obreb, area_m2, centroid_lon, centroid_lat) "
            "VALUES(?,?,?,?,?,?,?)",
            (id_rcn, "Testowo", f"ul. Testowa {id_rcn[-1]}", "0001", area, lon, lat))
    conn.commit()
    conn.close()
    return name


def test_list_plugins(client, auth, tmp_data_dir):
    _write_plugin(tmp_data_dir, "echo.py", ECHO_SRC)
    r = client.get("/api/plugins", auth=auth)
    assert r.status_code == 200
    items = {p["id"]: p for p in r.json()}
    assert items["echo"]["name"] == "Echo"
    assert items["echo"]["enabled"] is True and items["echo"]["error"] is None
    assert "run" not in items["echo"]


def test_run_plugin(client, auth, tmp_data_dir):
    _write_plugin(tmp_data_dir, "echo.py", ECHO_SRC)
    wid = _make_workspace_with_tx(tmp_data_dir)
    r = client.post(f"/api/workspaces/{wid}/plugins/echo", auth=auth,
                    json={"id_rcn_in": ["RCN-1", "RCN-2", "RCN-3", "RCN-4"]})
    assert r.status_code == 200, r.text
    out = r.json()
    assert out["text"] == "n=4 m2=3"
    assert out["table"]["has_poi"] is False  # brak pliku POI


def test_run_plugin_with_params(client, auth, tmp_data_dir):
    """body.params trafia do wtyczki przez ctx.params; brak params -> pusty dict."""
    _write_plugin(tmp_data_dir, "echo.py", ECHO_SRC)
    wid = _make_workspace_with_tx(tmp_data_dir)
    r = client.post(f"/api/workspaces/{wid}/plugins/echo", auth=auth,
                    json={"id_rcn_in": ["RCN-1"], "params": {"alpha": 0.15}})
    assert r.status_code == 200, r.text
    assert r.json()["table"]["alpha"] == 0.15
    r = client.post(f"/api/workspaces/{wid}/plugins/echo", auth=auth,
                    json={"id_rcn_in": ["RCN-1"]})
    assert r.status_code == 200
    assert r.json()["table"]["alpha"] is None


def test_run_plugin_capped_header(client, auth, tmp_data_dir):
    """Koszyk powyżej MAX_IDS -> nagłówek X-RCN-Capped (UI pokazuje notkę)."""
    from app.analysis import MAX_IDS
    _write_plugin(tmp_data_dir, "echo.py", ECHO_SRC)
    wid = _make_workspace_with_tx(tmp_data_dir)
    ids = ["RCN-1"] + [f"FAKE-{i}" for i in range(MAX_IDS)]  # MAX_IDS + 1 unikalnych
    r = client.post(f"/api/workspaces/{wid}/plugins/echo", auth=auth,
                    json={"id_rcn_in": ids})
    assert r.status_code == 200, r.text
    assert r.headers.get("X-RCN-Capped") == f"{MAX_IDS}/{MAX_IDS + 1}"
    # Poniżej limitu -- brak nagłówka.
    r = client.post(f"/api/workspaces/{wid}/plugins/echo", auth=auth,
                    json={"id_rcn_in": ["RCN-1"]})
    assert "X-RCN-Capped" not in r.headers


def test_run_unknown_plugin_404(client, auth, tmp_data_dir):
    wid = _make_workspace_with_tx(tmp_data_dir)
    r = client.post(f"/api/workspaces/{wid}/plugins/niema", auth=auth,
                    json={"id_rcn_in": ["RCN-1"]})
    assert r.status_code == 404


def test_run_broken_plugin_400(client, auth, tmp_data_dir):
    _write_plugin(tmp_data_dir, "zepsuta.py", "raise RuntimeError('boom')\n")
    wid = _make_workspace_with_tx(tmp_data_dir)
    r = client.post(f"/api/workspaces/{wid}/plugins/zepsuta", auth=auth,
                    json={"id_rcn_in": ["RCN-1"]})
    assert r.status_code == 400


def test_run_disabled_plugin_409(client, auth, tmp_data_dir):
    _write_plugin(tmp_data_dir, "echo.py.disabled", ECHO_SRC)
    wid = _make_workspace_with_tx(tmp_data_dir)
    r = client.post(f"/api/workspaces/{wid}/plugins/echo", auth=auth,
                    json={"id_rcn_in": ["RCN-1"]})
    assert r.status_code == 409


def test_run_export_kind_501(client, auth, tmp_data_dir):
    _write_plugin(tmp_data_dir, "exp.py", ECHO_SRC.replace(
        '"kind": "analysis"', '"kind": "export"').replace('"id": "echo"', '"id": "exp"'))
    wid = _make_workspace_with_tx(tmp_data_dir)
    r = client.post(f"/api/workspaces/{wid}/plugins/exp", auth=auth,
                    json={"id_rcn_in": ["RCN-1"]})
    assert r.status_code == 501


def test_run_plugin_exception_500(client, auth, tmp_data_dir):
    _write_plugin(tmp_data_dir, "rzuca.py", '''
PLUGIN = {"id": "rzuca", "name": "Rzuca", "kind": "analysis"}
def run(rows, ctx):
    raise ValueError("celowy wyjatek")
''')
    wid = _make_workspace_with_tx(tmp_data_dir)
    r = client.post(f"/api/workspaces/{wid}/plugins/rzuca", auth=auth,
                    json={"id_rcn_in": ["RCN-1"]})
    assert r.status_code == 500
    assert "ValueError" in r.json()["detail"]


def test_run_plugin_nan_500(client, auth, tmp_data_dir):
    _write_plugin(tmp_data_dir, "nan.py", '''
PLUGIN = {"id": "nan", "name": "NaN", "kind": "analysis"}
def run(rows, ctx):
    return {"x": float("nan")}
''')
    wid = _make_workspace_with_tx(tmp_data_dir)
    r = client.post(f"/api/workspaces/{wid}/plugins/nan", auth=auth,
                    json={"id_rcn_in": ["RCN-1"]})
    assert r.status_code == 500
    assert "NaN" in r.json()["detail"]


def test_upload_preview_then_confirm(client, auth, tmp_data_dir):
    files = {"file": ("dowolna-nazwa.py", ECHO_SRC.encode(), "text/x-python")}
    # Faza 1: preview -- plik NIE wylądował jeszcze w plugins/.
    r = client.post("/api/plugins/upload", auth=auth, files=files)
    assert r.status_code == 200, r.text
    assert r.json()["status"] == "preview"
    assert r.json()["plugin"]["id"] == "echo"
    assert not (tmp_data_dir / "plugins" / "echo.py").exists()
    # Faza 2: confirm -- zapis pod nazwą Z ID (nie z filename uploadu).
    r = client.post("/api/plugins/upload", auth=auth, files=files, data={"confirm": "true"})
    assert r.status_code == 200
    assert r.json()["status"] == "installed"
    assert (tmp_data_dir / "plugins" / "echo.py").exists()
    assert not (tmp_data_dir / "plugins" / "dowolna-nazwa.py").exists()
    # Kolizja id bez overwrite -> 409; z overwrite -> OK.
    r = client.post("/api/plugins/upload", auth=auth, files=files, data={"confirm": "true"})
    assert r.status_code == 409
    r = client.post("/api/plugins/upload", auth=auth, files=files,
                    data={"confirm": "true", "overwrite": "true"})
    assert r.status_code == 200


def test_upload_rejects_invalid(client, auth, tmp_data_dir):
    # Nie-.py -> 400.
    r = client.post("/api/plugins/upload", auth=auth,
                    files={"file": ("plik.txt", b"x", "text/plain")})
    assert r.status_code == 400
    # >1 MB -> 413.
    big = b"#" + b"x" * (1024 * 1024 + 10)
    r = client.post("/api/plugins/upload", auth=auth,
                    files={"file": ("big.py", big, "text/x-python")})
    assert r.status_code == 413
    # Zepsuta wtyczka -> 400 z powodem, nic nie zapisane.
    r = client.post("/api/plugins/upload", auth=auth,
                    files={"file": ("zla.py", b"raise RuntimeError('x')", "text/x-python")})
    assert r.status_code == 400
    assert "RuntimeError" in r.json()["detail"]
    leftover = list((tmp_data_dir / "plugins").iterdir()) \
        if (tmp_data_dir / "plugins").is_dir() else []
    assert leftover == []


def test_delete_and_enable(client, auth, tmp_data_dir):
    _write_plugin(tmp_data_dir, "echo.py", ECHO_SRC)
    # Wyłącz -> .py.disabled, na liście enabled=False.
    r = client.post("/api/plugins/echo/enable", auth=auth, json={"enabled": False})
    assert r.status_code == 200 and r.json()["enabled"] is False
    assert (tmp_data_dir / "plugins" / "echo.py.disabled").exists()
    # Włącz z powrotem.
    r = client.post("/api/plugins/echo/enable", auth=auth, json={"enabled": True})
    assert r.status_code == 200 and r.json()["error"] is None
    assert (tmp_data_dir / "plugins" / "echo.py").exists()
    # Usuń.
    r = client.delete("/api/plugins/echo", auth=auth)
    assert r.status_code == 200
    assert not (tmp_data_dir / "plugins" / "echo.py").exists()
    r = client.delete("/api/plugins/echo", auth=auth)
    assert r.status_code == 404


def test_delete_and_enable_when_filename_differs_from_id(client, auth, tmp_data_dir):
    """Wtyczka skopiowana ręcznie pod inną nazwą pliku niż PLUGIN['id'] --
    delete/enable mają działać po id (ścieżka z rejestru, nie z konwencji nazwy)."""
    _write_plugin(tmp_data_dir, "inna-nazwa.py", ECHO_SRC)  # PLUGIN['id'] == "echo"
    client.post("/api/plugins/rescan", auth=auth)
    # Wyłącz po id -> plik o INNEJ nazwie dostaje suffix .disabled.
    r = client.post("/api/plugins/echo/enable", auth=auth, json={"enabled": False})
    assert r.status_code == 200
    assert (tmp_data_dir / "plugins" / "inna-nazwa.py.disabled").exists()
    assert not (tmp_data_dir / "plugins" / "inna-nazwa.py").exists()
    # Po wyłączeniu rejestr zna wtyczkę pod id ze stemu pliku -- włącz i usuń po nim.
    r = client.post("/api/plugins/inna-nazwa/enable", auth=auth, json={"enabled": True})
    assert r.status_code == 200
    assert (tmp_data_dir / "plugins" / "inna-nazwa.py").exists()
    r = client.delete("/api/plugins/echo", auth=auth)
    assert r.status_code == 200
    assert not (tmp_data_dir / "plugins" / "inna-nazwa.py").exists()


def test_path_traversal_rejected(client, auth, tmp_data_dir):
    sentinel = tmp_data_dir / "evil.py"
    sentinel.write_text("x = 1")
    r = client.delete("/api/plugins/..%2Fevil", auth=auth)
    assert r.status_code in (400, 404)
    assert sentinel.exists()
    r = client.post("/api/plugins/..%2Fevil/enable", auth=auth, json={"enabled": False})
    assert r.status_code in (400, 404)
    assert sentinel.exists()


def test_rescan(client, auth, tmp_data_dir):
    r = client.post("/api/plugins/rescan", auth=auth)
    assert r.status_code == 200 and r.json() == []
    _write_plugin(tmp_data_dir, "echo.py", ECHO_SRC)
    r = client.post("/api/plugins/rescan", auth=auth)
    assert [p["id"] for p in r.json()] == ["echo"]


@pytest.fixture
def ro_client(tmp_data_dir, monkeypatch):
    """Client z dodatkowym kontem readonly (modyfikacje wtyczek -> 403)."""
    monkeypatch.setenv("RCN_READONLY_USER", "viewer")
    monkeypatch.setenv("RCN_READONLY_PASSWORD", "viewer-pass")
    for mod in list(sys.modules):
        if mod == "app.config" or mod.startswith("app.") or mod == "app":
            del sys.modules[mod]
    from fastapi.testclient import TestClient

    from app.main import app
    return TestClient(app)


def test_readonly_cannot_modify_plugins(ro_client, tmp_data_dir):
    ro = ("viewer", "viewer-pass")
    _write_plugin(tmp_data_dir, "echo.py", ECHO_SRC)
    # Odczyt + run dozwolone dla readonly.
    assert ro_client.get("/api/plugins", auth=ro).status_code == 200
    # Modyfikacje -> 403.
    r = ro_client.post("/api/plugins/upload", auth=ro,
                       files={"file": ("echo.py", ECHO_SRC.encode(), "text/x-python")})
    assert r.status_code == 403
    assert ro_client.delete("/api/plugins/echo", auth=ro).status_code == 403
    assert ro_client.post("/api/plugins/echo/enable", auth=ro,
                          json={"enabled": False}).status_code == 403


def test_export_plugin_result_xlsx(client, auth, tmp_data_dir):
    """Generyczny eksport wyniku wtyczki: tabele JSON -> skoroszyt (arkusz per
    tabela + Interpretacja); tabela dict k->v bez wiersza 'title'."""
    import io

    from openpyxl import load_workbook
    payload = {
        "filename": "Odległości analityczne",
        "tables": [
            {"title": "Podsumowanie", "columns": ["Kategoria", "Mediana [m]"],
             "rows": [["Szkoła", 600], ["Apteka", None]]},
            {"title": "Statystyki", "Średnia": 7187.44, "n": 316},
        ],
        "text": "Linia 1\nLinia 2",
    }
    r = client.post("/api/plugins/result.xlsx", auth=auth, json=payload)
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert 'filename="Odleg' in r.headers["content-disposition"] or \
        'filename="Odle' in r.headers["content-disposition"]

    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Podsumowanie", "Statystyki", "Interpretacja"]
    ws = wb["Podsumowanie"]
    assert [c.value for c in ws[1]] == ["Kategoria", "Mediana [m]"]
    assert ws.cell(row=2, column=1).value == "Szkoła" and ws.cell(row=2, column=2).value == 600
    assert ws.cell(row=3, column=2).value is None
    assert ws[1][0].font.bold is True
    ws2 = wb["Statystyki"]
    kv = {ws2.cell(row=i, column=1).value: ws2.cell(row=i, column=2).value
          for i in range(2, ws2.max_row + 1)}
    assert kv == {"Średnia": 7187.44, "n": 316}
    assert "title" not in kv
    assert wb["Interpretacja"]["A1"].value == "Linia 1"

    # Brak tabel i tekstu -> 400.
    r = client.post("/api/plugins/result.xlsx", auth=auth, json={"tables": []})
    assert r.status_code == 400


def test_analysis_regression_after_refactor(client, auth, tmp_data_dir):
    """Refaktor _load_rows nie zmienia kontraktu /analysis (klucze + wartości)."""
    wid = _make_workspace_with_tx(tmp_data_dir)
    r = client.post(f"/api/workspaces/{wid}/analysis", auth=auth,
                    json={"id_rcn_in": ["RCN-1", "RCN-2", "RCN-3", "RCN-4"]})
    assert r.status_code == 200, r.text
    out = r.json()
    assert set(out.keys()) == {
        "workspace_id", "total_requested", "total_matched", "capped",
        "stats_by_kind", "per_rynek", "per_rodzaj_transakcji", "top_obreb",
        "quarterly", "scatter",
    }
    assert out["total_matched"] == 4
    lokal = next(s for s in out["stats_by_kind"] if s["kind"] == "lokal")
    assert lokal["cena"]["n"] == 3
    assert lokal["cena"]["median"] == 450000.0
    assert lokal["cena_m2"]["n"] == 3
    # mediana cena/m2: 400000/50=8000, 450000/55=8181.82, 500000/60=8333.33
    assert lokal["cena_m2"]["median"] == 8181.82
    assert len(out["scatter"]) == 3  # tylko rekordy z area>0
    # Wszystkie 3 lokale w Q1 2025 -> jeden wpis kwartalny z n=3.
    assert len(out["quarterly"]) == 1
    assert out["quarterly"][0]["q"] == "2025-Q1" and out["quarterly"][0]["n"] == 3
